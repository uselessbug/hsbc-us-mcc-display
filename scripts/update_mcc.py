#!/usr/bin/env python3
"""Generate Mastercard MCC JSON from the official Quick Reference Booklet attachment."""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

PDF_URL = (
    "https://www.mastercard.com/content/dam/mccom/shared/business/support/"
    "rules-pdfs/mastercard-quick-reference-booklet-merchant.pdf"
)
MIN_ENTRIES = 300
MAX_ENTRIES = 1500
KNOWN = {
    "0742": ("veterinary",),
    "4121": ("taxi", "limousine"),
    "5812": ("restaurant", "eating"),
    "7011": ("hotel", "motel", "lodging"),
}


def run(*args: str, cwd: Path | None = None) -> str:
    return subprocess.run(
        args, cwd=cwd, check=True, stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT, text=True,
    ).stdout


def download(url: str, path: Path) -> None:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (compatible; HSBC-MCC-DB-Updater/1.0)",
            "Accept": "application/pdf,*/*;q=0.8",
        },
    )
    with urllib.request.urlopen(req, timeout=60) as response:  # noqa: S310
        if response.status != 200:
            raise RuntimeError(f"Download failed: HTTP {response.status}")
        data = response.read()
    if not data.startswith(b"%PDF-"):
        raise RuntimeError("Mastercard source did not return a PDF")
    path.write_bytes(data)


def attachments(pdf: Path, out: Path) -> list[Path]:
    if not shutil.which("pdfdetach"):
        raise RuntimeError("pdfdetach is required (install poppler-utils)")
    out.mkdir(parents=True, exist_ok=True)
    listing = run("pdfdetach", "-list", str(pdf))
    if ".xlsx" not in listing.lower() and ".xlsm" not in listing.lower():
        raise RuntimeError("No Excel attachment advertised by Mastercard PDF")
    run("pdfdetach", "-saveall", str(pdf), cwd=out)
    return sorted(p for p in out.iterdir() if p.suffix.lower() in {".xlsx", ".xlsm"})


def workbook(files: list[Path]) -> Path:
    if not files:
        raise RuntimeError("No Excel attachment found")
    def score(p: Path) -> tuple[int, str]:
        n = p.name.lower()
        return (
            4 * ("mcc" in n) + 2 * ("merchant" in n)
            + 2 * ("listing" in n) + ("comprehensive" in n),
            n,
        )
    files.sort(key=score, reverse=True)
    return files[0]


def code(value: object) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return f"{value:04d}" if 0 <= value <= 9999 else None
    if isinstance(value, float) and value.is_integer():
        i = int(value)
        return f"{i:04d}" if 0 <= i <= 9999 else None
    s = str(value).strip()
    return s.zfill(4) if re.fullmatch(r"\d{1,4}", s) else None


def description(value: object) -> str | None:
    if value is None:
        return None
    s = re.sub(r"\s+", " ", str(value)).strip()
    return s if len(s) >= 3 else None


def header(rows: list[tuple[object, ...]]) -> tuple[int, int, int] | None:
    for row_no, row in enumerate(rows[:50]):
        cells = [re.sub(r"\s+", " ", str(v or "")).strip().lower() for v in row]
        mcc = [
            i for i, v in enumerate(cells)
            if v == "mcc" or "merchant category code" in v or "acceptor business code" in v
        ]
        desc = [
            i for i, v in enumerate(cells)
            if "description" in v or "merchant category" in v or "acceptor business" in v
        ]
        desc = [i for i in desc if i not in mcc]
        if mcc and desc:
            return row_no, mcc[0], desc[0]
    return None


def sheet_entries(sheet) -> dict[str, str] | None:
    h = header(list(sheet.iter_rows(min_row=1, max_row=50, values_only=True)))
    if not h:
        return None
    row_no, mcc_col, desc_col = h
    result: dict[str, str] = {}
    for row in sheet.iter_rows(min_row=row_no + 2, values_only=True):
        if max(mcc_col, desc_col) >= len(row):
            continue
        c = code(row[mcc_col])
        d = description(row[desc_col])
        if not c or not d:
            continue
        previous = result.get(c)
        if previous and previous.casefold() != d.casefold():
            raise RuntimeError(f"Conflicting descriptions for MCC {c}: {previous!r} vs {d!r}")
        result[c] = d
    return dict(sorted(result.items(), key=lambda item: int(item[0])))


def extract(path: Path) -> tuple[dict[str, str], str]:
    book = load_workbook(path, read_only=True, data_only=True)
    candidates: list[tuple[int, int, str, dict[str, str]]] = []
    for sheet in book.worksheets:
        entries = sheet_entries(sheet)
        if not entries:
            continue
        name = sheet.title.lower()
        name_score = 5 * ("mcc" in name) + 4 * ("merchant category" in name) + 3 * ("listing" in name)
        candidates.append((len(entries), name_score, sheet.title, entries))
    if not candidates:
        raise RuntimeError("No recognizable MCC worksheet found")
    candidates.sort(key=lambda x: (x[0], x[1]), reverse=True)
    count, _, title, entries = candidates[0]
    print(f"Selected worksheet {title!r} with {count} entries")
    return entries, title


def document_date(pdf: Path) -> str:
    if not shutil.which("pdftotext"):
        raise RuntimeError("pdftotext is required (install poppler-utils)")
    with tempfile.TemporaryDirectory() as tmp:
        text_file = Path(tmp) / "pages.txt"
        run("pdftotext", "-f", "1", "-l", "8", str(pdf), str(text_file))
        text = text_file.read_text(encoding="utf-8", errors="replace")
    match = re.search(
        r"\b(\d{1,2})\s+"
        r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+"
        r"(20\d{2})\b",
        text,
        re.IGNORECASE,
    )
    if not match:
        raise RuntimeError("Unable to determine Mastercard document date")
    return dt.datetime.strptime(" ".join(match.groups()), "%d %B %Y").date().isoformat()


def validate(entries: dict[str, str]) -> None:
    if not MIN_ENTRIES <= len(entries) <= MAX_ENTRIES:
        raise RuntimeError(f"Unexpected MCC count: {len(entries)}")
    for c, words in KNOWN.items():
        d = entries.get(c, "").lower()
        if not d or not any(word in d for word in words):
            raise RuntimeError(f"Known MCC {c} missing or unexpected: {entries.get(c)!r}")


def write(entries: dict[str, str], date: str, sheet: str, output: Path) -> None:
    payload = {
        "schemaVersion": 1,
        "version": date,
        "source": {
            "network": "Mastercard",
            "document": "Quick Reference Booklet - Merchant Edition",
            "url": PDF_URL,
            "documentDate": date,
            "worksheet": sheet,
        },
        "mcc": entries,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("data/mcc-mastercard.json"))
    args = parser.parse_args()

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        pdf = root / "mastercard.pdf"
        print(f"Downloading {PDF_URL}")
        download(PDF_URL, pdf)
        date = document_date(pdf)
        print(f"Document date: {date}")
        xlsx = workbook(attachments(pdf, root / "attachments"))
        print(f"Workbook: {xlsx.name}")
        entries, sheet = extract(xlsx)
        validate(entries)
        print(f"Validated {len(entries)} MCC entries")
        write(entries, date, sheet, args.output)
        print(f"Wrote {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
